import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


def clean(text):
    return " ".join((text or "").replace("\u00a0", " ").split())


def xml_text(element):
    return clean("".join(node.text or "" for node in element.iter(qn("w:t"))))


def iter_blocks(parent):
    body = parent.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def table_rows(table):
    rows = []
    for row in table.rows:
        values = []
        seen = set()
        for cell in row.cells:
            cell_id = id(cell._tc)
            if cell_id in seen:
                continue
            seen.add(cell_id)
            values.append(clean(cell.text))
        rows.append(values)
    return rows


def main(path):
    doc = Document(path)
    blocks = []
    rendered_page = 1
    for index, block in enumerate(iter_blocks(doc), start=1):
        xml = block._element
        if isinstance(block, Paragraph):
            text = xml_text(xml)
            style = block.style.name if block.style is not None else ""
            blocks.append({
                "index": index,
                "type": "paragraph",
                "page_hint": rendered_page,
                "style": style,
                "text": text,
            })
        else:
            blocks.append({
                "index": index,
                "type": "table",
                "page_hint": rendered_page,
                "rows": table_rows(block),
            })
        page_breaks = sum(1 for node in xml.iter(qn("w:lastRenderedPageBreak")))
        explicit_breaks = sum(
            1
            for node in xml.iter(qn("w:br"))
            if node.get(qn("w:type")) == "page"
        )
        rendered_page += page_breaks + explicit_breaks

    headers = []
    footers = []
    for section_index, section in enumerate(doc.sections, start=1):
        headers.append({
            "section": section_index,
            "text": clean(" ".join(p.text for p in section.header.paragraphs)),
        })
        footers.append({
            "section": section_index,
            "text": clean(" ".join(p.text for p in section.footer.paragraphs)),
        })

    media = []
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            media.append(str(rel.target_ref))

    result = {
        "path": str(Path(path)),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "inline_shape_count": len(doc.inline_shapes),
        "section_count": len(doc.sections),
        "media": media,
        "headers": headers,
        "footers": footers,
        "blocks": blocks,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main(sys.argv[1])

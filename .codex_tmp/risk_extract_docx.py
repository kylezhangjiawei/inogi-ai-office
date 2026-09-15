from pathlib import Path
import sys
import argparse
from docx import Document
from docx.document import Document as DocumentType
from docx.table import Table
from docx.text.paragraph import Paragraph


SOURCE = Path(
    r"C:\Users\Administrator\Documents\xwechat_files\zjw19920831_8887\msg\file\2026-09\采购与生产部门AI应用建议及落地手册_V1.0.docx"
)

sys.stdout.reconfigure(encoding="utf-8")

parser = argparse.ArgumentParser()
parser.add_argument("--start", type=int, default=1)
parser.add_argument("--end", type=int, default=10_000)
args = parser.parse_args()


def iter_blocks(parent):
    if isinstance(parent, DocumentType):
        parent_element = parent.element.body
    else:
        raise TypeError(type(parent))
    for child in parent_element.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, parent)
        elif child.tag.endswith("}tbl"):
            yield Table(child, parent)


doc = Document(SOURCE)
print("# BODY")
for index, block in enumerate(iter_blocks(doc), start=1):
    if index < args.start or index > args.end:
        continue
    if isinstance(block, Paragraph):
        text = " ".join(block.text.split())
        if text:
            print(f"P{index:04d}\t[{block.style.name}]\t{text}")
    else:
        print(f"T{index:04d}\t{len(block.rows)}x{len(block.columns)}")
        for row_index, row in enumerate(block.rows, start=1):
            cells = [" / ".join(" ".join(p.text.split()) for p in cell.paragraphs if p.text.strip()) for cell in row.cells]
            print(f"  R{row_index:03d}\t" + "\t".join(cells))

for section_index, section in enumerate(doc.sections, start=1):
    for label, container in (("HEADER", section.header), ("FOOTER", section.footer)):
        print(f"# {label} {section_index}")
        for paragraph in container.paragraphs:
            text = " ".join(paragraph.text.split())
            if text:
                print(f"[{paragraph.style.name}]\t{text}")
